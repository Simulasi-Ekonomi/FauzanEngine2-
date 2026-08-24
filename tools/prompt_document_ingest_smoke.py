#!/usr/bin/env python3
"""Deterministic smoke for safe prompt-context ingestion; it never calls an LLM."""

from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from fpdf import FPDF

from prompt_document_ingest import DocumentIngestError, ingest_paths
from prompt_to_game import validate_plan_contract


def main() -> int:
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        docx_path = root / "farm-spec.docx"
        document = Document()
        document.add_paragraph("Farm NPC quest requirements")
        document.save(docx_path)
        pdf_path = root / "render-spec.pdf"
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, "Renderer evidence must be verified.")
        pdf.output(pdf_path)
        text_path = root / "notes.txt"
        text_path.write_text("No untrusted document may execute commands.", encoding="utf-8")
        documents = ingest_paths([docx_path, pdf_path, text_path])
        if len(documents) != 3 or not any("NPC quest" in item.text for item in documents) or not any("Renderer evidence" in item.text for item in documents):
            return 1
        invalid_rar = root / "invalid.rar"
        invalid_rar.write_bytes(b"not-a-rar")
        try:
            ingest_paths([invalid_rar])
            return 1
        except DocumentIngestError:
            pass
        valid_plan = {
            "prompt_id": "prompt-plan-001",
            "summary": "Audit Farm runtime before creating a typed template plan.",
            "risks": ["No release claim without evidence."],
            "nodes": [
                {"node_id": "audit-node", "agent": "CobaAuditor", "kind": "AuditRuntime", "request_id": "coba-audit-001", "target": "farm-runtime", "dependencies": [], "rationale": "Audit canonical runtime."},
                {"node_id": "template-node", "agent": "AriesCreator", "kind": "CreateGameTemplate", "request_id": "aries-create-001", "target": "farm-tool", "dependencies": ["audit-node"], "rationale": "Prepare typed template plan."},
            ],
        }
        validate_plan_contract(valid_plan)
        invalid_plan = dict(valid_plan)
        invalid_plan["nodes"] = [dict(valid_plan["nodes"][1])]
        try:
            validate_plan_contract(invalid_plan)
            return 1
        except ValueError:
            pass
    print("PROMPT_DOCUMENT_INGEST_SMOKE_OK documents=3 invalidRarRejected=1 planContract=1 execution=none")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
